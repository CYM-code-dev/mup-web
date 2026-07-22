# -*- coding: utf-8 -*-
"""
多目标物 Excel 模板读写 (测量数据版)

①②③ 方法级共享参数 + ④ 标准品分组(结构/类型/目标物名/分组共享标准参数/各物纯度·证书) 由网页填写;
本模板只承载**测量数据**: 目标物表(各物 曲线方法/内标物/标液浓度/加标体积) + 精密度&回收率表(每物多组平行样 实测加标C/称样量m) + 曲线表(浓度,峰面积)。
上传时页面标准参数 ∪ 本模板测量数据 → 合并出完整报告。

每目标物可独立选 外标法/内标法 + 标注内标物(如"内标物A"); 由页面 ⑤ 设置预填进模板, 上传照读。

模板工作表:
  目标物          [目标物, 分组, 曲线方法, (内标物), 标液浓度(mg/L), 加标体积(μL)]。名+分组+曲线方法+(内标物) 预填; 「内标物」列仅当存在内标法目标物时才有(全外标→无此列); 标液浓度/加标体积 待填 (X/曲线p/x_pred 不再手填, 合并段由⑥实测加标C派生)
  精密度&回收率   长表 [目标物, 实测加标C(mg/L), 称样量m(g)]。每物多组平行样(按页面⑥测定次数预填 N 行/物 仅名, C/m 待填); 系统按 各物标液浓度·加标体积 + 方法V 派生 R=C/C₀, w=C·V/m; 每物≥2 对。多次进样(>1)→末列加「进样序号」
  曲线拟合        长表 [目标物, 浓度mg/L, 分析物峰面积(, 内标峰面积)]。存在内标法目标物时才加「内标峰面积」列; 外标 y=分析物峰面积, 内标 y=分析物峰面积/内标峰面积(系统算)。按页面曲线点数预填 N 行/目标物(仅填名), 每物≥3 点。多次进样(>1)→末列加「进样序号」
"""
import openpyxl
import unicodedata
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
BOLD = Font(bold=True)


def _disp_width(s):
    """字符串显示宽度: CJK/全角=2, 其余=1 (Excel 列宽单位≈一个数字宽)。"""
    return sum(2 if unicodedata.east_asian_width(c) in ("W", "F") else 1 for c in str(s))


def _header(ws, hdr):
    for c, h in enumerate(hdr, 1):
        cell = ws.cell(1, c, h)
        cell.font = BOLD
        cell.fill = HEADER_FILL
        ws.column_dimensions[get_column_letter(c)].width = _disp_width(h) + 2


def _long_sheet(wb, name, hdr, rows):
    ws = wb.create_sheet(name)
    _header(ws, hdr)
    for i, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1):
            ws.cell(i, c, v)


def _group_long(ws, arity):
    """长表 → {目标物: [值或元组]}。arity=每个目标物的数据列数。"""
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        name = str(row[0]).strip()
        data = list(row[1:1 + arity])
        if len(data) < arity or any(v is None for v in data):
            continue
        if arity == 1:
            out.setdefault(name, []).append(float(data[0]))
        else:
            out.setdefault(name, []).append(tuple(float(v) for v in data))
    return out


def _to_float(v):
    """单元格 → float; 空/非法 → None。"""
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


_INJ_LABELS = {1: "一次进样", 2: "二次进样", 3: "三次进样", 4: "四次进样", 5: "五次进样",
               6: "六次进样", 7: "七次进样", 8: "八次进样", 9: "九次进样", 10: "十次进样"}


def _inj_label(n):
    """进样序号列标签: 1→一次进样 … 10→十次进样; 超出→第N次进样。"""
    return _INJ_LABELS.get(n) or f"第{n}次进样"


def _read_pr_sheet(ws):
    """合并表 [目标物, 测定值, 回收率] → {nm: {"replicates":[...], "recovery":[...]}}。
    两值列各自可选(空/非法跳过); 允许重复性与回收率行数不同。"""
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        nm = str(row[0]).strip()
        if not nm:
            continue
        d = out.setdefault(nm, {"replicates": [], "recovery": []})
        rep = _to_float(row[1]) if len(row) > 1 else None
        rec = _to_float(row[2]) if len(row) > 2 else None
        if rep is not None:
            d["replicates"].append(rep)
        if rec is not None:
            d["recovery"].append(rec)
    return out


def _read_spike_sheet(ws):
    """合并表 [目标物, 实测加标C(mg/L), 称样量m(g)] → {nm: {"spike_C":[...], "spike_m":[...]}}。
    仅当一行 C、m 均合法才计为一对 (对齐单模式⑥: 缺一不可); R/w 由合并段派生。"""
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        nm = str(row[0]).strip()
        if not nm:
            continue
        c = _to_float(row[1]) if len(row) > 1 else None
        m = _to_float(row[2]) if len(row) > 2 else None
        if c is None or m is None:
            continue
        d = out.setdefault(nm, {"spike_C": [], "spike_m": []})
        d["spike_C"].append(c)
        d["spike_m"].append(m)
    return out


def _read_curve_sheet(ws, methods):
    """新格式曲线长表 [目标物, 浓度mg/L, 分析物峰面积, 内标峰面积] →
    {目标物: [(浓度, y), ...]}。按目标物 method 算 y:
      内标法 → 分析物峰面积/内标峰面积 (缺/0/非法 → 记 missing, 跳过该点);
      外标法 → 分析物峰面积 (内标峰面积列忽略)。
    methods: {目标物: "外标法"/"内标法"}。返回 (out, missing_set)。"""
    out, missing = {}, set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        nm = str(row[0]).strip()
        if not nm:
            continue
        conc = _to_float(row[1]) if len(row) > 1 else None
        a_area = _to_float(row[2]) if len(row) > 2 else None
        if conc is None or a_area is None:
            continue
        if methods.get(nm) == "内标法":
            ia = _to_float(row[3]) if len(row) > 3 else None
            if ia is None or ia == 0:
                missing.add(nm)
                continue
            y = a_area / ia
        else:
            y = a_area
        out.setdefault(nm, []).append((conc, y))
    return out, missing


def _num(d, key, cast):
    """单元格 → cast(float/int); 空/非法 → None。"""
    v = d.get(key)
    if v is None or str(v).strip() == "":
        return None
    try:
        return cast(v)
    except (TypeError, ValueError):
        return None


def _first_num(rows, key, cast):
    """合并表多行(同名目标物)中取该列首个合法数; 全空 → None。"""
    for d in rows:
        v = _num(d, key, cast)
        if v is not None:
            return v
    return None


def write_template(path, groups=None, curve_meta=None, n_points=3, n_reps=2, n_inj_point=1, n_inj_meas=1):
    """生成「仅测量数据」xlsx 模板。标准品参数(纯度/证书/分组共享)在页面填, 不在本模板。
    groups: [{name, kind, analytes:[名]}] 结构 (来自页面 _rows_to_groups); None→1 个 demo 分组。
    curve_meta: {名: {"method":"外标法"/"内标法", "is_name":内标物标签}}; 预填 目标物表
    「曲线方法」「内标物」两列。None/缺→默认外标法、内标物空。
    n_points: 每目标物曲线点数; 据此预填「曲线」表 N 行/目标物(仅目标物名, 浓度/峰面积留空待填)。
    n_reps: 每目标物加标平行样次数; 据此预填「精密度&回收率」表 N 行/目标物(仅目标物名, C/m 留空待填), ≥2。
    n_inj_point: 每点标液进样次数; 「曲线」表行数/物 = n_points × n_inj_point (同浓度多行不同峰面积), ≥1。
    n_inj_meas: 每次测定进样次数; 「精密度&回收率」表行数/物 = n_reps × n_inj_meas, ≥1。
    多次进样(>1)时两表末列加「进样序号」(一次/二次…, 点内分组); =1 时无此列。"""
    if groups is None:
        groups = [{"name": "纯品-1", "kind": "solid", "analytes": []}]
    curve_meta = curve_meta or {}
    n_points = max(3, int(n_points or 3))
    n_reps = max(2, int(n_reps or 2))
    n_inj_point = max(1, int(n_inj_point or 1))
    n_inj_meas = max(1, int(n_inj_meas or 1))
    # ponytail: 「内标物」列(目标物表)与「内标峰面积」列(曲线表)仅当存在内标法目标物时才加; 全外标→目标物5列/曲线3列。读侧 read_template 按表头自适应。
    has_is = any((curve_meta.get(nm, {}).get("method") == "内标法")
                 for g in groups for nm in g.get("analytes", []))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "目标物"
    _header(ws, ["目标物", "分组", "曲线方法"] + (["内标物"] if has_is else [])
            + ["标液浓度(mg/L)", "加标体积(μL)"])
    r = 2
    for g in groups:
        for nm in g.get("analytes", []):
            cm = curve_meta.get(nm, {})
            ws.cell(r, 1, nm)
            ws.cell(r, 2, g["name"])
            ws.cell(r, 3, cm.get("method") or "外标法")
            if has_is:
                ws.cell(r, 4, cm.get("is_name") or "")
            r += 1

    curve_hdr = ["目标物", "浓度mg/L", "分析物峰面积"] + (["内标峰面积"] if has_is else [])
    # 曲线表预填 N 行/目标物(仅填目标物名); 浓度/峰面积由用户补。N = n_points × n_inj_point (每点重复进样)。
    # ⑥ 加标数据(实测加标C/称样量m)单列「精密度&回收率」表: 每物多组平行样(按 n_reps × n_inj_meas 预填 N 行/物 仅名, C/m 待填); 每物≥2 对
    # 多次进样(>1)→末列加「进样序号」(一次/二次…, 点内分组); 读侧按列序读前几列, 末列标签被忽略 → 向后兼容。
    pr_hdr = ["目标物", "实测加标C(mg/L)", "称样量m(g)"]
    if n_inj_meas > 1:
        pr_hdr.append("进样序号")
        _pad = len(pr_hdr) - 2   # 目标物与末列「进样序号」之间的数据列数 (C/m)
        pr_rows = [[nm] + [None] * _pad + [_inj_label(j)] for g in groups for nm in g.get("analytes", [])
                   for _ in range(n_reps) for j in range(1, n_inj_meas + 1)]
    else:
        pr_rows = [(nm,) for g in groups for nm in g.get("analytes", []) for _ in range(n_reps)]
    _long_sheet(wb, "精密度&回收率", pr_hdr, pr_rows)
    if n_inj_point > 1:
        curve_hdr.append("进样序号")
        _cpad = len(curve_hdr) - 2   # 目标物与「进样序号」之间: 浓度/分析物峰面积(±内标峰面积)
        curve_rows = [[nm] + [None] * _cpad + [_inj_label(j)] for g in groups for nm in g.get("analytes", [])
                      for _ in range(n_points) for j in range(1, n_inj_point + 1)]
    else:
        curve_rows = [(nm,) for g in groups for nm in g.get("analytes", []) for _ in range(n_points)]
    _long_sheet(wb, "曲线拟合", curve_hdr, curve_rows)
    # 目标物列(A)宽度以最长化合物名为准 (名长>表头时撑开; 其余列仍按表头宽)
    _a_w = max(_disp_width("目标物"),
               max((_disp_width(nm) for g in groups for nm in g.get("analytes", [])), default=0)) + 2
    for _ws in wb.worksheets:
        _ws.column_dimensions["A"].width = _a_w
    wb.save(path)


def read_template(path):
    """解析测量数据模板 → {目标物名: {X, p, x_pred, curve_method, is_name, points, replicates, recovery}}。
    标准品参数在页面填, 不在本模板。每目标物可独立 外标法/内标法 (目标物表「曲线方法」「内标物」列);
    曲线表新格式 [浓度, 分析物峰面积, 内标峰面积] → 内标法 y=面积比, 外标法 y=分析物峰面积。
    兼容旧模板: 目标物表无「曲线方法」列→默认外标; 曲线表旧 3 列 [浓度, y]→y 直用; 旧版分「重复性」「回收率」两表→分别读。
    目标物表 [目标物, 分组, 曲线方法, (内标物), 标液浓度(mg/L), 加标体积(μL)] (X/曲线p/x_pred 不再手填; 合并段由⑥实测加标C 派生 p=个数、x_pred=均值, X 走测定值w均值兜底)。「内标物」列仅当存在内标法目标物时才有; 全外标模板无此列→is_name=""。
    兼容旧模板: 上版 目标物表含 X_mgkg/曲线p/x_pred_mgL 三列 → 照读(无⑥加标C时回退用)。
    ⑥ 加标数据(实测加标C/称样量m) 单列「精密度&回收率」表 [目标物, 实测加标C(mg/L), 称样量m(g)], 每物多组平行样(多行)。
    合并段按 各物标液浓度·加标体积 + 方法V 派生 R=C/C₀, w=C·V/m。兼容: 上版 目标物表末两列并C/m / 更旧 精密度&回收率 测定值格式 / 最旧 分两表。
    校验: 每个目标物 ≥3 曲线点; 加标C/m≥2 对 (或旧 重复·回收各≥2); 内标法曲线点须有内标峰面积, 否则 raise ValueError。"""
    wb = openpyxl.load_workbook(path, data_only=True)
    if "目标物" not in wb.sheetnames:
        raise ValueError("模板无「目标物」表: 请在页面 ④ 重新下载模板。")
    ws = wb["目标物"]
    headers = [str(c.value).strip() for c in ws[1]]
    has_pr_in_target = any("实测加标" in h for h in headers)   # 合并格式: ⑥ C/m 在目标物表末两列
    raw = {}   # nm -> [行 dict, ...] (合并表同名目标物可多行: 首行各物参数 + 各重复行 C/m)
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        d = dict(zip(headers, row))
        nm = str(d.get("目标物", "")).strip()
        if nm:
            raw.setdefault(nm, []).append(d)
    out = {}
    for nm, rows in raw.items():
        cm = "外标法"
        for d in rows:
            vs = str(d.get("曲线方法")).strip() if d.get("曲线方法") is not None else ""
            if vs in ("外标法", "内标法"):
                cm = vs
                break
        isn = ""
        for d in rows:
            v = d.get("内标物")
            if v is not None and str(v).strip():
                isn = str(v).strip()
                break
        spike_C, spike_m = [], []
        if has_pr_in_target:   # 合并格式: 收集各行 (C,m) 配对
            for d in rows:
                c = _num(d, "实测加标C(mg/L)", float)
                m = _num(d, "称样量m(g)", float)
                if c is not None and m is not None:
                    spike_C.append(c)
                    spike_m.append(m)
        out[nm] = {"X": _first_num(rows, "X_mgkg", float), "p": _first_num(rows, "曲线p", int),
                   "x_pred": _first_num(rows, "x_pred_mgL", float),
                   "curve_method": cm, "is_name": isn,
                   "points": [], "replicates": [], "recovery": [],
                   "spike_std_conc": _first_num(rows, "标液浓度(mg/L)", float),
                   "spike_add_vol": _first_num(rows, "加标体积(μL)", float),
                   "spike_C": spike_C, "spike_m": spike_m}
    methods = {nm: a["curve_method"] for nm, a in out.items()}
    # 曲线: 新格式(分析物峰面积/内标峰面积)按 method 算 y; 旧格式(单列 y)直读
    _curve_sn = next((s for s in ("曲线拟合", "曲线") if s in wb.sheetnames), None)   # 优先「曲线拟合」, 兼容旧「曲线」
    curve_ws = wb[_curve_sn] if _curve_sn else None
    curve_hdr = ([str(c.value).strip() for c in curve_ws[1]] if curve_ws else [])
    errors = []
    if any(h and "内标峰面积" in h for h in curve_hdr):
        curves, missing_is = _read_curve_sheet(curve_ws, methods)
        for nm in sorted(missing_is):
            errors.append(f"{nm}: 内标法曲线点缺/非法内标峰面积")
    else:
        curves = _group_long(curve_ws, 2) if curve_ws else {}
    # ⑥ 加标C/m: 合并格式已在目标物表读好; 旧模板回退 精密度&回收率 表 / 更旧 分两表
    pr_ws = wb["精密度&回收率"] if "精密度&回收率" in wb.sheetnames else None
    pr_hdr = ([str(c.value).strip() for c in pr_ws[1]] if pr_ws else [])
    reps, recs = {}, {}
    pr_spike_fmt = has_pr_in_target
    if not has_pr_in_target:
        if pr_ws and any("实测加标" in h for h in pr_hdr):
            for nm, sp in _read_spike_sheet(pr_ws).items():
                if nm in out:
                    out[nm]["spike_C"] = sp["spike_C"]
                    out[nm]["spike_m"] = sp["spike_m"]
            pr_spike_fmt = True
        elif pr_ws:   # 上上版格式: 直接 测定值/回收率
            pr = _read_pr_sheet(pr_ws)
            reps = {nm: d["replicates"] for nm, d in pr.items()}
            recs = {nm: d["recovery"] for nm, d in pr.items()}
        else:         # 更旧: 分开的 重复性/回收率 两表
            reps = _group_long(wb["重复性"], 1) if "重复性" in wb.sheetnames else {}
            recs = _group_long(wb["回收率"], 1) if "回收率" in wb.sheetnames else {}
    for nm, a in out.items():
        a["points"] = curves.get(nm, [])
        if not has_pr_in_target:   # 合并格式: reps/recs 留空待合并段派生; 旧格式从分表填
            a["replicates"] = reps.get(nm, [])
            a["recovery"] = recs.get(nm, [])
        if len(a["points"]) < 3:
            errors.append(f"{nm}: 曲线点<3")
        if pr_spike_fmt:
            if len(a["spike_C"]) < 2:
                errors.append(f"{nm}: 加标C/m配对<2")
        else:
            if len(a["replicates"]) < 2:
                errors.append(f"{nm}: 重复性<2")
            if len(a["recovery"]) < 2:
                errors.append(f"{nm}: 回收率<2")
    if errors:
        raise ValueError("模板解析错误:\n- " + "\n- ".join(errors))
    return out


def legacy_to_groups(stock, work_uses, u_work, analytes):
    """旧 (方法级 stock, work_uses, u_work, analytes[]) → (groups, analytes) 分组模型。
    solid → 每物一固体分组; liquid_dilute → 一液体分组含全员。供旧草稿(v1)迁移。
    analytes 旧键: purity/U_purity/k_purity/m_std/stock_flask_volume/X/p/x_pred/(Urel_cert/k_cert)。"""
    src = (stock or {}).get("stock_source", "solid")
    work_uses = list(work_uses or [])
    has_work = bool(work_uses)
    groups, out = [], []
    if src == "liquid_dilute":
        g = {"name": "液标", "kind": "liquid", "analytes": [a["name"] for a in analytes],
             "work_uses": work_uses, "cert_mode": "relative",
             "Urel_cert": stock.get("Urel_cert"), "k_cert": stock.get("k_cert"),
             "pip_kind": stock.get("pip_kind"), "pip_vol": stock.get("pip_vol"),
             "flask_vol": stock.get("flask_vol")}
        if not has_work and u_work is not None:
            g["u_work"] = u_work
        groups.append(g)
        for a in analytes:
            na = {"name": a["name"], "group": "液标", "X": a.get("X"), "p": a.get("p"),
                  "x_pred": a.get("x_pred"), "points": a.get("points", []),
                  "replicates": a.get("replicates", []), "recovery": a.get("recovery", [])}
            if a.get("Urel_cert") is not None:
                na["Urel_cert"] = a["Urel_cert"]
            if a.get("k_cert") is not None:
                na["k_cert"] = a["k_cert"]
            out.append(na)
    else:
        for a in analytes:
            nm = a["name"]
            g = {"name": f"纯品-{nm}", "kind": "solid", "analytes": [nm],
                 "m_std_g": a.get("m_std"), "n_weighings": a.get("n_weighings", 2),
                 "flask_volume_mL": a.get("stock_flask_volume"), "work_uses": work_uses}
            if not has_work and u_work is not None:
                g["u_work"] = u_work
            groups.append(g)
            out.append({"name": nm, "group": g["name"], "purity": a.get("purity"),
                        "U_purity": a.get("U_purity"), "k_purity": a.get("k_purity"),
                        "X": a.get("X"), "p": a.get("p"), "x_pred": a.get("x_pred"),
                        "points": a.get("points", []), "replicates": a.get("replicates", []),
                        "recovery": a.get("recovery", [])})
    return groups, out


if __name__ == "__main__":
    import os
    out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "_tmpl.xlsx")
    groups = [{"name": "混标", "kind": "solid", "analytes": ["A_外标", "B_内标"]}]
    curve_meta = {"A_外标": {"method": "外标法", "is_name": ""},
                  "B_内标": {"method": "内标法", "is_name": "内标物A"}}
    write_template(out, groups, curve_meta, n_points=3)
    wb = openpyxl.load_workbook(out)
    assert "标准品分组" not in wb.sheetnames, "不应有标准品分组 sheet"
    assert "说明" not in wb.sheetnames and wb.sheetnames == ["目标物", "精密度&回收率", "曲线拟合"], wb.sheetnames
    # 目标物表: 曲线方法/内标物 已预填; 标液浓度/加标体积 留空 (6 列; X/曲线p/x_pred 已删, 合并段派生)
    rows = list(wb["目标物"].iter_rows(min_row=2, values_only=True))
    assert rows[:2] == [("A_外标", "混标", "外标法", None, None, None),
                        ("B_内标", "混标", "内标法", "内标物A", None, None)], rows[:2]
    # 精密度&回收率表: [目标物, 实测加标C(mg/L), 称样量m(g)]; 预填 2 行/目标物(仅名, C/m 待填)
    assert [str(c.value).strip() for c in wb["精密度&回收率"][1]] == \
        ["目标物", "实测加标C(mg/L)", "称样量m(g)"]
    # 列宽按表头自适应 (CJK=2宽) → 列名完整展示
    assert wb["精密度&回收率"].column_dimensions["B"].width == _disp_width("实测加标C(mg/L)") + 2
    # 目标物列(A)宽度以最长化合物名为准: 名长>表头 → 三表 A 列均撑开
    _outL = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "_tmpl_long.xlsx")
    write_template(_outL, [{"name": "G", "kind": "solid", "analytes": ["六邻苯二甲酸酯类化合物"]}], {}, n_points=3)
    _wbL = openpyxl.load_workbook(_outL)
    assert [s.column_dimensions["A"].width for s in _wbL.worksheets] == \
        [_disp_width("六邻苯二甲酸酯类化合物") + 2] * 3
    os.remove(_outL)
    pr_names = [r[0] for r in wb["精密度&回收率"].iter_rows(min_row=2, values_only=True)]
    assert pr_names == ["A_外标", "A_外标", "B_内标", "B_内标"], pr_names
    # 曲线表头为 4 列
    assert [str(c.value).strip() for c in wb["曲线拟合"][1]] == \
        ["目标物", "浓度mg/L", "分析物峰面积", "内标峰面积"]
    # 曲线表按 n_points 预填 (仅目标物名, 浓度/峰面积留空): A_外标×3 + B_内标×3
    cv_names = [r[0] for r in wb["曲线拟合"].iter_rows(min_row=2, values_only=True)]
    assert cv_names == ["A_外标"] * 3 + ["B_内标"] * 3, cv_names
    # 模拟用户填测量数据: 目标物表每物 标液浓度/加标体积 (X/p/x_pred 不再手填, 合并段派生); 加标C/m 走 精密度&回收率表(每物2对 平行样)
    ws = wb["目标物"]
    ws.cell(2, 5, 100.0); ws.cell(2, 6, 100)                        # A_外标 标液100mg/L·100μL
    ws.cell(3, 5, 200.0); ws.cell(3, 6, 100)                        # B_内标 标液200mg/L·100μL
    pr = wb["精密度&回收率"]                                          # 加标C/m 平行样(每物2对)
    pr.cell(2, 2, 12.34); pr.cell(2, 3, 0.5004)                     # A_外标 第1对
    pr.cell(3, 2, 12.10); pr.cell(3, 3, 0.4998)                     # A_外标 第2对
    pr.cell(4, 2, 21.80); pr.cell(4, 3, 0.5001)                     # B_内标 第1对
    pr.cell(5, 2, 21.50); pr.cell(5, 3, 0.4999)                     # B_内标 第2对
    cv = wb["曲线拟合"]
    # A_外标: 外标法, 第4列(内标峰面积)留空
    for i, (n, x, a) in enumerate([("A_外标", 1, 100), ("A_外标", 2, 200), ("A_外标", 3, 300)], start=2):
        cv.cell(i, 1, n); cv.cell(i, 2, x); cv.cell(i, 3, a)
    # B_内标: 内标法, 填 分析物峰面积 + 内标峰面积
    for i, (n, x, a, isa) in enumerate(
            [("B_内标", 1, 100, 50), ("B_内标", 2, 200, 50), ("B_内标", 3, 300, 50)], start=5):
        cv.cell(i, 1, n); cv.cell(i, 2, x); cv.cell(i, 3, a); cv.cell(i, 4, isa)
    wb.save(out)
    meas = read_template(out)
    assert set(meas) == {"A_外标", "B_内标"}, meas
    # 外标: y = 分析物峰面积; method/is_name 回读
    assert meas["A_外标"]["curve_method"] == "外标法" and meas["A_外标"]["is_name"] == ""
    assert meas["A_外标"]["points"] == [(1.0, 100.0), (2.0, 200.0), (3.0, 300.0)], meas["A_外标"]["points"]
    # 内标: y = 分析物峰面积/内标峰面积 (=2,4,6)
    assert meas["B_内标"]["curve_method"] == "内标法" and meas["B_内标"]["is_name"] == "内标物A"
    assert meas["B_内标"]["points"] == [(1.0, 2.0), (2.0, 4.0), (3.0, 6.0)], meas["B_内标"]["points"]
    # 精密度&回收率表 → 加标C/m 配对; 目标物表 → 各物标液参数; X/p/x_pred 列已删→None (合并段从spike_C派生); reps/recs 留空待合并段派生
    assert meas["A_外标"]["X"] is None and meas["A_外标"]["p"] is None and meas["A_外标"]["x_pred"] is None
    assert meas["A_外标"]["spike_C"] == [12.34, 12.10], meas["A_外标"]["spike_C"]
    assert meas["A_外标"]["spike_m"] == [0.5004, 0.4998]
    assert meas["A_外标"]["spike_std_conc"] == 100.0 and meas["A_外标"]["spike_add_vol"] == 100.0
    assert meas["B_内标"]["spike_C"] == [21.80, 21.50]
    assert meas["B_内标"]["spike_std_conc"] == 200.0 and meas["B_内标"]["spike_add_vol"] == 100.0
    assert meas["A_外标"]["replicates"] == [] and meas["A_外标"]["recovery"] == []

    # n_reps 控制精密度&回收率预填行数/物 (默认 2; 这里验 3)
    out_r = out + ".nreps.xlsx"
    write_template(out_r, groups, curve_meta, n_points=3, n_reps=3)
    _pr_r = [r[0] for r in openpyxl.load_workbook(out_r)["精密度&回收率"].iter_rows(min_row=2, values_only=True)]
    assert _pr_r == ["A_外标"] * 3 + ["B_内标"] * 3, _pr_r
    os.remove(out_r)

    # n_inj_point/n_inj_meas 倍增行数 + 末列「进样序号」(默认 1 无此列; 这里验 2: 曲线 3点×2=6/物, 精密度 2对×2=4/物)
    out_i = out + ".ninj.xlsx"
    write_template(out_i, groups, curve_meta, n_points=3, n_reps=2, n_inj_point=2, n_inj_meas=2)
    _wb_i = openpyxl.load_workbook(out_i)
    _pr_i = list(_wb_i["精密度&回收率"].iter_rows(min_row=2, values_only=True))
    assert [r[0] for r in _pr_i] == ["A_外标"] * 4 + ["B_内标"] * 4
    assert [str(c.value).strip() for c in _wb_i["精密度&回收率"][1]] == \
        ["目标物", "实测加标C(mg/L)", "称样量m(g)", "进样序号"]
    # 点内分组: 每对 一次/二次 → A 四行 = 一次,二次,一次,二次
    assert [r[3] for r in _pr_i] == ["一次进样", "二次进样"] * 4, [r[3] for r in _pr_i]
    _cv_i = list(_wb_i["曲线拟合"].iter_rows(min_row=2, values_only=True))
    assert [r[0] for r in _cv_i] == ["A_外标"] * 6 + ["B_内标"] * 6
    # 曲线 4 列(含内标峰面积) + 进样序号 → 5 列; 每点 一次/二次 → A 六行 = (一次,二次)×3
    assert [str(c.value).strip() for c in _wb_i["曲线拟合"][1]] == \
        ["目标物", "浓度mg/L", "分析物峰面积", "内标峰面积", "进样序号"]
    assert [r[4] for r in _cv_i] == ["一次进样", "二次进样"] * 6, [r[4] for r in _cv_i]
    # 读侧按列序读前几列, 末列「进样序号」被忽略 → 填重复浓度也不丢点 (6 点全读出)
    _cv = _wb_i["曲线拟合"]
    for i, (x, a) in enumerate([(1, 10), (1, 11), (2, 20), (2, 21), (3, 30), (3, 31)], start=2):
        _cv.cell(i, 2, x); _cv.cell(i, 3, a)                       # A_外标 6 行 (外标法, 第4列留空)
    for i, (x, a, isa) in enumerate([(1, 100, 50), (1, 110, 50), (2, 200, 50), (2, 210, 50), (3, 300, 50), (3, 310, 50)],
                                    start=8):
        _cv.cell(i, 2, x); _cv.cell(i, 3, a); _cv.cell(i, 4, isa)  # B_内标 6 行
    # 目标物表 标液浓度/加标体积 (合并段派生 C₀); 精密度&回收率表 填 C/m (末列进样序号已预填, 读侧忽略)
    _tg = _wb_i["目标物"]
    _tg.cell(2, 5, 100.0); _tg.cell(2, 6, 100); _tg.cell(3, 5, 200.0); _tg.cell(3, 6, 100)
    _pr_w = _wb_i["精密度&回收率"]
    for i, (c, m) in enumerate([(12.34, 0.5004), (12.10, 0.4998), (12.20, 0.5001), (12.05, 0.4999)], start=2):
        _pr_w.cell(i, 2, c); _pr_w.cell(i, 3, m)                  # A_外标 4 行
    for i, (c, m) in enumerate([(21.80, 0.5001), (21.50, 0.4999), (21.65, 0.5002), (21.45, 0.4998)], start=6):
        _pr_w.cell(i, 2, c); _pr_w.cell(i, 3, m)                  # B_内标 4 行
    _wb_i.save(out_i)
    _meas_i = read_template(out_i)
    assert _meas_i["A_外标"]["points"] == [(1, 10), (1, 11), (2, 20), (2, 21), (3, 30), (3, 31)], _meas_i["A_外标"]["points"]
    assert _meas_i["B_内标"]["points"] == [(1, 2.0), (1, 2.2), (2, 4.0), (2, 4.2), (3, 6.0), (3, 6.2)], _meas_i["B_内标"]["points"]
    assert _meas_i["A_外标"]["spike_C"] == [12.34, 12.10, 12.20, 12.05], _meas_i["A_外标"]["spike_C"]
    os.remove(out_i)

    # 全外标 → 目标物表 5 列(无「内标物」); 曲线表 3 列 (无「内标峰面积」); 含内标 → 各多一列
    out2 = out + ".all_ext.xlsx"
    _g2 = [{"name": "混标", "kind": "solid", "analytes": ["A", "B"]}]
    write_template(out2, _g2, {"A": {"method": "外标法", "is_name": ""},
                              "B": {"method": "外标法", "is_name": ""}})
    wb2 = openpyxl.load_workbook(out2)
    assert [str(c.value).strip() for c in wb2["目标物"][1]] == \
        ["目标物", "分组", "曲线方法", "标液浓度(mg/L)", "加标体积(μL)"], [str(c.value) for c in wb2["目标物"][1]]
    assert [str(c.value).strip() for c in wb2["曲线拟合"][1]] == \
        ["目标物", "浓度mg/L", "分析物峰面积"], [str(c.value) for c in wb2["曲线拟合"][1]]
    os.remove(out2)

    os.remove(out)
    print("OK 测量数据模板写读自洽 (无说明sheet; 目标物含内标6列/全外标5列 X/p/x_pred删→合并段派生; 外标+内标; 全外标曲线3列; 曲线预填n_points; ⑥加标C/m独立表+n_reps行; 进样次数倍增曲线/精密度行)")
